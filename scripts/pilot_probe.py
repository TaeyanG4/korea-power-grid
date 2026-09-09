from __future__ import annotations
import json, zipfile
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
SOURCES={'demand':ROOT/'data/raw/demand/2026-07/demand_2026_07.zip','dispatch':ROOT/'data/raw/dispatch/2026-07/dispatch_2026_07.zip','state_estimation':ROOT/'data/raw/state_estimation/2026-07/state_estimation_2026_07.zip'}
def recover(name,utf8):
    if utf8: return name
    try: return name.encode('cp437').decode('cp949')
    except Exception: return name
def detect(data):
    sample=data[:2097152]
    for enc in ('utf-8-sig','utf-8','cp949','euc-kr'):
        try: sample.decode(enc); return enc
        except UnicodeDecodeError: pass
    return None
def probe_text(data):
    enc=detect(data)
    if not enc: return {'encoding':None}
    lines=data[:2097152].decode(enc).splitlines()[:5]
    return {'encoding':enc,'preview':lines,'delimiter_counts_first_5_lines':{'tab':sum(x.count('\t') for x in lines),'comma':sum(x.count(',') for x in lines),'pipe':sum(x.count('|') for x in lines),'semicolon':sum(x.count(';') for x in lines)}}
def probe_xlsx(data):
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        preview=[]
        for i,row in enumerate(ws.iter_rows(values_only=True)):
            preview.append(list(row))
            if i>=4: break
        sheets.append({'title':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'preview':preview})
    return {'sheets':sheets}
out={'month':'2026-07','sources':{}}
for src,path in SOURCES.items():
    r={'zip_path':str(path.relative_to(ROOT)).replace('\\','/'),'zip_size_bytes':path.stat().st_size}
    with zipfile.ZipFile(path) as z:
        r['zip_test']=z.testzip(); ms=[]
        for info in z.infolist():
            name=recover(info.filename,bool(info.flag_bits & 0x800)); data=z.read(info.filename)
            p=probe_xlsx(data) if name.lower().endswith('.xlsx') else probe_text(data)
            ms.append({'raw_zip_name':info.filename,'recovered_name':name,'size_bytes':info.file_size,'compressed_size_bytes':info.compress_size,'probe':p})
        r['members']=ms
    out['sources'][src]=r
p=ROOT/'data/audits/pilot_probe_2026_07.json'; p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2)); print('WROTE',p)
