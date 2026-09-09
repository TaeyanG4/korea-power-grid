from __future__ import annotations

import codecs
import csv
import hashlib
import io
import itertools
import json
import math
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MONTH = "2026-07"
EXPECTED = pd.date_range("2026-07-01 00:00:00", "2026-07-31 23:55:00", freq="5min")
OUT_DIR = ROOT / "data" / "processed" / "pilot" / MONTH
AUDIT_PATH = ROOT / "data" / "audits" / "pilot_2026_07.json"
OUT_DIR.mkdir(parents=True, exist_ok=True)
AUDIT_PATH.parent.mkdir(parents=True, exist_ok=True)

SOURCES = {
    "demand": {
        "zip": ROOT / "data/raw/demand/2026-07/demand_2026_07.zip",
        "article_url": "https://www.kpx.or.kr/board.es?act=view&bid=0065&list_no=78033&mid=a10109020700&nPage=1&tag=",
        "attachment_url": "https://www.kpx.or.kr/boardDownload.es?bid=0065&list_no=78033&seq=1",
    },
    "dispatch": {
        "zip": ROOT / "data/raw/dispatch/2026-07/dispatch_2026_07.zip",
        "article_url": "https://www.kpx.or.kr/board.es?act=view&bid=0070&list_no=78025&mid=a10109020200&nPage=1&tag=",
        "attachment_url": "https://www.kpx.or.kr/boardDownload.es?bid=0070&list_no=78025&seq=1",
    },
    "state_estimation": {
        "zip": ROOT / "data/raw/state_estimation/2026-07/state_estimation_2026_07.zip",
        "article_url": "https://www.kpx.or.kr/board.es?act=view&bid=0068&list_no=78031&mid=a10109020400&nPage=1&tag=",
        "attachment_url": "https://www.kpx.or.kr/boardDownload.es?bid=0068&list_no=78031&seq=1",
    },
}

TEXT_HEADER_SCAN_LIMIT = 100


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def recover_name(info: zipfile.ZipInfo) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def detect_text_encoding(probe: bytes) -> str:
    if probe.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        probe.decode("ascii")
        return "ascii"
    except UnicodeDecodeError:
        pass

    for encoding in ("utf-8", "cp949"):
        try:
            decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
            decoder.decode(probe, final=False)
            return encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError(
        "kpx-text",
        probe,
        0,
        min(len(probe), 1),
        "probe is neither ASCII, UTF-8, nor CP949",
    )


def _normalize_header_token(value: str) -> str:
    return "".join(value.strip().lstrip("\ufeff").upper().split())


def _is_three_column_header(source: str, columns: list[str]) -> bool:
    if len(columns) != 3:
        return False

    first, second, third = (_normalize_header_token(value) for value in columns)
    first_ok = first in {"TIME", "\uc2dc\uac04"}
    second_ok = (
        ("GEN" in second or "\ubc1c\uc804\uae30" in second)
        and ("CODE" in second or "\ucf54\ub4dc" in second)
    )
    if not (first_ok and second_ok):
        return False

    if source == "dispatch":
        return third == "BASEPOINT"
    if source == "state_estimation":
        return "MW" in third and ("\uc0c1\ud0dc\ucd94\uc815" in third or "EST" in third)
    return False


def inspect_three_column_text_header(
    source: str,
    stream,
    *,
    label: str = "<stream>",
    max_lines: int = TEXT_HEADER_SCAN_LIMIT,
) -> tuple[str, list[str], int]:
    probe = stream.read(4096)
    encoding = detect_text_encoding(probe)
    stream.seek(0)

    for index in range(max_lines):
        raw_line = stream.readline()
        if not raw_line:
            break
        decoded = raw_line.decode(encoding).strip()
        columns = [part.strip() for part in next(csv.reader([decoded]))]
        if _is_three_column_header(source, columns):
            return encoding, columns, index

    raise RuntimeError(
        f"{source}: could not locate expected 3-column CSV header in first "
        f"{max_lines} lines of {label}"
    )


def _dispatch_xlsx_header_columns(row: tuple) -> tuple[int, int, int] | None:
    values = ["" if value is None else _normalize_header_token(str(value)) for value in row]
    for start in range(max(0, len(values) - 2)):
        first, second, third = values[start : start + 3]
        if third == "BASEPOINT" and "CODE" in second and first:
            return start, start + 1, start + 2
    return None


def _is_xlsx_timestamp_value(value) -> bool:
    if isinstance(value, (datetime, pd.Timestamp)):
        return True
    if not isinstance(value, str):
        return False
    text = value.strip()
    if len(text) < 10 or "-" not in text:
        return False
    return bool(text[:4].isdigit() and text[5:7].isdigit() and text[8:10].isdigit())


def _dispatch_xlsx_data_columns(row: tuple) -> tuple[int, int, int] | None:
    for start in range(max(0, len(row) - 2)):
        first, second, third = row[start : start + 3]
        if _is_xlsx_timestamp_value(first) and second is not None and third is not None:
            return start, start + 1, start + 2
    return None


def read_dispatch_xlsx(data: bytes) -> tuple[pd.DataFrame, dict]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    frames: list[pd.DataFrame] = []
    layouts: list[dict] = []

    for worksheet in workbook.worksheets:
        row_iter = worksheet.iter_rows(values_only=True)
        probe_rows = list(itertools.islice(row_iter, 20))
        header_index = None
        columns = None

        for index, row in enumerate(probe_rows):
            columns = _dispatch_xlsx_header_columns(row)
            if columns is not None:
                header_index = index
                break

        if columns is not None:
            data_start = header_index + 1
        else:
            data_start = None
            for index, row in enumerate(probe_rows):
                columns = _dispatch_xlsx_data_columns(row)
                if columns is not None:
                    data_start = index
                    break

        if columns is None or data_start is None:
            raise RuntimeError(
                f"dispatch xlsx: could not locate a 3-column data block in worksheet "
                f"{worksheet.title!r}"
            )

        c0, c1, c2 = columns
        all_rows = itertools.chain(probe_rows[data_start:], row_iter)

        def selected_rows():
            for row in all_rows:
                if len(row) <= c2:
                    continue
                values = (row[c0], row[c1], row[c2])
                if all(value is None for value in values):
                    continue
                if not _is_xlsx_timestamp_value(values[0]):
                    continue
                yield values

        frame = pd.DataFrame(
            selected_rows(),
            columns=["TIME", "GEN_CODE", "BASEPOINT"],
        )
        frames.append(frame)
        layouts.append(
            {
                "worksheet": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "header_row_1based": None if header_index is None else header_index + 1,
                "data_columns_1based": [c0 + 1, c1 + 1, c2 + 1],
                "parsed_rows": int(len(frame)),
            }
        )

    workbook.close()
    if not frames:
        raise RuntimeError("dispatch xlsx: workbook has no readable worksheets")

    return pd.concat(frames, ignore_index=True), {
        "worksheet_count": len(layouts),
        "worksheet_layout": layouts,
    }


def select_source_zip_member(
    source: str,
    archive: zipfile.ZipFile,
) -> tuple[zipfile.ZipInfo, list[str]]:
    """Select the one archive member matching the requested source schema."""
    infos = archive.infolist()
    if not infos:
        raise RuntimeError(f"{source}: ZIP archive has no members")
    if len(infos) == 1:
        return infos[0], []

    candidates: list[zipfile.ZipInfo] = []
    for info in infos:
        try:
            with archive.open(info) as handle:
                probe = handle.read(65536)

            if source == "demand":
                if probe[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
                    candidates.append(info)
                    continue
                if probe[:4] == b"PK\x03\x04" and Path(info.filename).suffix.lower() in {
                    ".xlsx",
                    ".xlsm",
                }:
                    candidates.append(info)
                    continue
                encoding = detect_text_encoding(probe[:4096])
                text = probe.decode(encoding)
                if any(
                    "," in line
                    and ("시간" in line or line.lstrip("\ufeff").upper().startswith("TIME,"))
                    for line in text.splitlines()[:20]
                ):
                    candidates.append(info)
                continue

            if source == "dispatch" and probe[:4] == b"PK\x03\x04":
                if Path(info.filename).suffix.lower() in {".xlsx", ".xlsm"}:
                    candidates.append(info)
                continue

            if probe[:4] == b"PK\x03\x04":
                continue
            with archive.open(info) as handle:
                inspect_three_column_text_header(
                    source,
                    handle,
                    label=info.filename,
                )
            candidates.append(info)
        except (RuntimeError, UnicodeDecodeError, csv.Error):
            continue

    if len(candidates) != 1:
        raise RuntimeError(
            f"{source}: expected exactly one source-matching ZIP member, "
            f"found {len(candidates)} among {len(infos)} members"
        )

    selected = candidates[0]
    ignored = [recover_name(info) for info in infos if info is not selected]
    return selected, ignored


def finite_number_summary(series: pd.Series) -> dict:
    num = pd.to_numeric(series, errors="coerce")
    values = num.to_numpy(dtype="float64", na_value=np.nan)
    finite = values[np.isfinite(values)]
    q = {}
    if finite.size:
        for p in (0.0, 0.001, 0.01, 0.5, 0.99, 0.999, 1.0):
            q[str(p)] = float(np.quantile(finite, p))
    return {
        "null_or_parse_failure_count": int(num.isna().sum()),
        "nan_count": int(np.isnan(values).sum()),
        "pos_inf_count": int(np.isposinf(values).sum()),
        "neg_inf_count": int(np.isneginf(values).sum()),
        "negative_count": int((finite < 0).sum()),
        "zero_count": int((finite == 0).sum()),
        "quantiles": q,
    }


def parse_timestamp_series(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return pd.to_datetime(series, errors="coerce")
    raw = series.astype("string").str.strip().str.replace("/", "-", regex=False)
    parsed = pd.to_datetime(raw, errors="coerce", format="mixed")
    missing = parsed.isna() & raw.notna()
    if missing.any():
        normalized = (
            raw.loc[missing]
            .str.replace(" \uc624\uc804 ", " AM ", regex=False)
            .str.replace(" \uc624\ud6c4 ", " PM ", regex=False)
        )
        parsed.loc[missing] = pd.to_datetime(
            normalized,
            errors="coerce",
            format="%Y-%m-%d %p %I:%M:%S",
        )
    return parsed

def timestamp_summary(ts: pd.Series) -> dict:
    parsed = pd.to_datetime(ts, errors="coerce", format="mixed")
    valid = parsed.dropna()
    unique = pd.DatetimeIndex(valid.unique()).sort_values()
    expected_set = set(EXPECTED.to_pydatetime())
    unique_set = set(unique.to_pydatetime())
    missing = sorted(expected_set - unique_set)
    outside = sorted(unique_set - expected_set)
    deltas = unique.to_series().diff().dropna()
    delta_counts = {str(k): int(v) for k, v in deltas.value_counts().head(10).items()}
    return {
        "parse_failure_count": int(parsed.isna().sum()),
        "min": None if valid.empty else valid.min().isoformat(),
        "max": None if valid.empty else valid.max().isoformat(),
        "unique_timestamp_count": int(len(unique)),
        "expected_timestamp_count": int(len(EXPECTED)),
        "missing_expected_timestamp_count": int(len(missing)),
        "missing_expected_timestamp_examples": [x.isoformat() for x in missing[:20]],
        "outside_expected_timestamp_count": int(len(outside)),
        "outside_expected_timestamp_examples": [x.isoformat() for x in outside[:20]],
        "observed_delta_counts_top10": delta_counts,
        "timezone": "unverified; source timestamps stored as naive values in pilot",
    }


def diff_summary(a: pd.Series, b: pd.Series) -> dict:
    joined = pd.concat([a.rename("a"), b.rename("b")], axis=1).dropna()
    if joined.empty:
        return {"matched_timestamp_count": 0}
    d = joined["a"] - joined["b"]
    return {
        "matched_timestamp_count": int(len(joined)),
        "mean_difference_mw": float(d.mean()),
        "median_difference_mw": float(d.median()),
        "mean_absolute_difference_mw": float(d.abs().mean()),
        "difference_quantiles_mw": {
            "0.01": float(d.quantile(0.01)),
            "0.5": float(d.quantile(0.5)),
            "0.99": float(d.quantile(0.99)),
        },
    }


def read_source(source: str, path: Path) -> tuple[pd.DataFrame, dict]:
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        infos = z.infolist()
        info, ignored_members = select_source_zip_member(source, z)
        name = recover_name(info)
        raw_name = info.filename
        if source == "demand":
            data = z.read(info.filename)
            magic = data[:8].hex()
            if data[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
                df = pd.read_excel(io.BytesIO(data), engine="xlrd")
                physical_format = "Excel 97-2003 OLE .xls binary despite .xlsx member filename"
                reader = "pandas.read_excel(engine='xlrd')"
                encoding = None
                preamble_rows = 0
            elif data[:4] == b"PK\x03\x04":
                df = pd.read_excel(io.BytesIO(data), engine="openpyxl")
                physical_format = "Office Open XML .xlsx"
                reader = "pandas.read_excel(engine='openpyxl')"
                encoding = None
                preamble_rows = 0
            else:
                encoding = detect_text_encoding(data[:4096])
                text = data.decode(encoding)
                lines = text.splitlines()
                header_index = next(
                    (
                        idx
                        for idx, line in enumerate(lines[:20])
                        if "," in line and ("\uc2dc\uac04" in line or line.upper().startswith("TIME,"))
                    ),
                    None,
                )
                if header_index is None:
                    raise RuntimeError(
                        f"Could not find demand CSV header in first 20 lines; magic={magic}"
                    )
                df = pd.read_csv(io.StringIO("\n".join(lines[header_index:])))
                physical_format = f"{encoding} comma-delimited text with optional preamble"
                reader = f"pandas.read_csv(decoded {encoding} text)"
                preamble_rows = header_index
            physical = {
                "zip_test_bad_member": bad,
                "outer_zip_member_count": len(infos),
                "ignored_outer_zip_members": ignored_members,
                "raw_zip_member_name": raw_name,
                "recovered_member_name": name,
                "member_size_bytes": info.file_size,
                "member_compressed_size_bytes": info.compress_size,
                "member_magic_hex": magic,
                "physical_format": physical_format,
                "reader": reader,
                "encoding_used_for_full_parse": encoding,
                "preamble_rows_skipped": preamble_rows,
            }
        else:
            with z.open(info) as f:
                magic = f.read(8)

            if source == "dispatch" and magic[:4] == b"PK\x03\x04":
                data = z.read(info)
                df, workbook_meta = read_dispatch_xlsx(data)
                physical = {
                    "zip_test_bad_member": bad,
                    "outer_zip_member_count": len(infos),
                    "ignored_outer_zip_members": ignored_members,
                    "raw_zip_member_name": raw_name,
                    "recovered_member_name": name,
                    "member_size_bytes": info.file_size,
                    "member_compressed_size_bytes": info.compress_size,
                    "member_magic_hex": magic.hex(),
                    "physical_format": "Office Open XML .xlsx workbook",
                    "reader": "openpyxl.load_workbook(read_only=True, data_only=True)",
                    "encoding_used_for_full_parse": None,
                    "preamble_rows_skipped": None,
                    **workbook_meta,
                }
            else:
                with z.open(info) as f:
                    encoding, columns, preamble_rows = inspect_three_column_text_header(
                        source,
                        f,
                        label=name,
                    )
                    df = pd.read_csv(
                        f,
                        encoding=encoding,
                        dtype=str,
                        names=columns,
                        header=None,
                        low_memory=False,
                    )
                physical = {
                    "zip_test_bad_member": bad,
                    "outer_zip_member_count": len(infos),
                    "ignored_outer_zip_members": ignored_members,
                    "raw_zip_member_name": raw_name,
                    "recovered_member_name": name,
                    "member_size_bytes": info.file_size,
                    "member_compressed_size_bytes": info.compress_size,
                    "physical_format": "comma-delimited text",
                    "encoding_used_for_full_parse": encoding,
                    "preamble_rows_skipped": preamble_rows,
                }
    return df, physical


def normalize(source: str, df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    original_columns = [str(c) for c in df.columns]
    if source == "demand":
        if len(original_columns) != 2:
            raise RuntimeError(f"demand expected 2 columns, got {original_columns}")
        out = pd.DataFrame({
            "timestamp": parse_timestamp_series(df.iloc[:, 0]),
            "demand_forecast_mw": pd.to_numeric(df.iloc[:, 1], errors="coerce"),
        })
        value_col = "demand_forecast_mw"
        grain = ["timestamp"]
        mapping = {original_columns[0]: "timestamp", original_columns[1]: value_col}
    else:
        if len(original_columns) != 3:
            raise RuntimeError(f"{source} expected 3 columns, got {original_columns}")
        value_col = "dispatch_mw" if source == "dispatch" else "estimated_generation_mw"
        out = pd.DataFrame({
            "timestamp": parse_timestamp_series(df.iloc[:, 0]),
            "generator_id": df.iloc[:, 1].astype("string").str.strip(),
            value_col: pd.to_numeric(df.iloc[:, 2], errors="coerce"),
        })
        grain = ["timestamp", "generator_id"]
        mapping = {
            original_columns[0]: "timestamp",
            original_columns[1]: "generator_id",
            original_columns[2]: value_col,
        }
    return out, {
        "original_columns": original_columns,
        "canonical_mapping_for_pilot": mapping,
        "candidate_grain": grain,
        "value_column": value_col,
    }


def source_audit(source: str, cfg: dict) -> tuple[dict, set[str], pd.Series | None]:
    started = time.perf_counter()
    path: Path = cfg["zip"]
    read_started = time.perf_counter()
    raw, physical = read_source(source, path)
    read_seconds = time.perf_counter() - read_started
    canonical, schema = normalize(source, raw)
    value_col = schema["value_column"]
    ts_info = timestamp_summary(canonical["timestamp"])
    duplicate_count = int(canonical.duplicated(schema["candidate_grain"], keep=False).sum())
    duplicate_group_count = int(canonical.duplicated(schema["candidate_grain"], keep=False).groupby(canonical[schema["candidate_grain"]].apply(tuple, axis=1)).any().sum()) if False else None

    generators: set[str] = set()
    generator_null_count = None
    if "generator_id" in canonical.columns:
        empty = canonical["generator_id"].isna() | canonical["generator_id"].eq("")
        generator_null_count = int(empty.sum())
        generators = set(canonical.loc[~empty, "generator_id"].astype(str).unique())

    parquet_path = OUT_DIR / f"{source}.parquet"
    write_started = time.perf_counter()
    canonical.to_parquet(parquet_path, index=False, compression="zstd")
    write_seconds = time.perf_counter() - write_started

    aggregate = None
    if source in ("dispatch", "state_estimation"):
        aggregate = canonical.groupby("timestamp", observed=True)[value_col].sum(min_count=1)
    elif source == "demand":
        aggregate = canonical.drop_duplicates("timestamp").set_index("timestamp")[value_col]

    member_total = sum(m["member_size_bytes"] for m in [physical])
    stat = path.stat()
    audit = {
        "source": source,
        "month": MONTH,
        "source_article_url": cfg["article_url"],
        "attachment_url": cfg["attachment_url"],
        "local_zip_path": str(path.relative_to(ROOT)).replace("\\", "/"),
        "zip_size_bytes": stat.st_size,
        "sha256": sha256(path),
        "downloaded_at_utc_from_mtime": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        "status": "success",
        "retry_count": 0,
        "physical": physical,
        "schema": schema,
        "row_count": int(len(canonical)),
        "timestamp": ts_info,
        "duplicate_candidate_key_row_count": duplicate_count,
        "generator_id_null_or_blank_count": generator_null_count,
        "unique_generator_count": None if source == "demand" else int(len(generators)),
        "numeric": finite_number_summary(canonical[value_col]),
        "parquet": {
            "path": str(parquet_path.relative_to(ROOT)).replace("\\", "/"),
            "compression": "zstd",
            "size_bytes": parquet_path.stat().st_size,
            "row_count": int(len(canonical)),
        },
        "timing_seconds": {
            "read_parse": round(read_seconds, 6),
            "parquet_write": round(write_seconds, 6),
            "total_source": round(time.perf_counter() - started, 6),
        },
    }
    del raw, canonical
    return audit, generators, aggregate


def projection(total_zip: int, total_parquet: int) -> dict:
    horizons = {"3_years": 36, "5_years": 60, "8_years": 96, "10_years": 120, "full_132_month_board_history": 132}
    out = {}
    for label, months in horizons.items():
        out[label] = {
            "months": months,
            "raw_zip_linear_bytes": total_zip * months,
            "raw_zip_linear_gb_decimal": round(total_zip * months / 1_000_000_000, 3),
            "parquet_linear_bytes": total_parquet * months,
            "parquet_linear_gb_decimal": round(total_parquet * months / 1_000_000_000, 3),
        }
    return out


def main() -> None:
    total_started = time.perf_counter()
    audit = {
        "pilot_month": MONTH,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "expected_5_minute_timestamps": len(EXPECTED),
        "sources": {},
        "cross_source": {},
        "projection": {},
        "notes": [
            "Timezone is not asserted by this pilot because no explicit timezone metadata has yet been verified from the source files.",
            "Linear size projections use only the 2026-07 pilot and must be replaced by measured checkpoint sizes during historical backfill.",
            "Cross-source aggregate differences are descriptive metrics, not hard QA equality rules.",
        ],
    }
    generator_sets = {}
    aggregates = {}
    for source, cfg in SOURCES.items():
        print(f"AUDITING {source}...", flush=True)
        result, generators, aggregate = source_audit(source, cfg)
        audit["sources"][source] = result
        generator_sets[source] = generators
        aggregates[source] = aggregate
        print(
            f"DONE {source}: rows={result['row_count']:,}, parquet={result['parquet']['size_bytes']:,} bytes",
            flush=True,
        )

    dispatch_ids = generator_sets["dispatch"]
    state_ids = generator_sets["state_estimation"]
    overlap = dispatch_ids & state_ids
    union = dispatch_ids | state_ids
    audit["cross_source"]["generator_id_overlap"] = {
        "dispatch_unique": len(dispatch_ids),
        "state_estimation_unique": len(state_ids),
        "overlap": len(overlap),
        "dispatch_only": len(dispatch_ids - state_ids),
        "state_only": len(state_ids - dispatch_ids),
        "jaccard": None if not union else round(len(overlap) / len(union), 8),
    }
    audit["cross_source"]["aggregate_comparisons"] = {
        "dispatch_sum_minus_demand_forecast": diff_summary(aggregates["dispatch"], aggregates["demand"]),
        "state_sum_minus_demand_forecast": diff_summary(aggregates["state_estimation"], aggregates["demand"]),
        "state_sum_minus_dispatch_sum": diff_summary(aggregates["state_estimation"], aggregates["dispatch"]),
    }

    total_zip = sum(x["zip_size_bytes"] for x in audit["sources"].values())
    total_extracted = sum(x["physical"]["member_size_bytes"] for x in audit["sources"].values())
    total_parquet = sum(x["parquet"]["size_bytes"] for x in audit["sources"].values())
    audit["pilot_totals"] = {
        "zip_size_bytes": total_zip,
        "extracted_member_size_bytes": total_extracted,
        "parquet_size_bytes": total_parquet,
        "parquet_vs_zip_ratio": round(total_parquet / total_zip, 6),
        "parquet_vs_extracted_ratio": round(total_parquet / total_extracted, 6),
    }
    audit["projection"] = projection(total_zip, total_parquet)
    audit["timing_seconds_total"] = round(time.perf_counter() - total_started, 6)
    AUDIT_PATH.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"WROTE {AUDIT_PATH}")
    print(json.dumps({
        "pilot_totals": audit["pilot_totals"],
        "projection": audit["projection"],
        "cross_source": audit["cross_source"],
        "timing_seconds_total": audit["timing_seconds_total"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()



